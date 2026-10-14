"""
Travel Accommodation Portal (internal network)
Flask + pyodbc (SQL Server) + MSAL/Microsoft Graph mail + signed approval links
"""
import base64
import html
import io
import ipaddress
import json
import logging
import os
import re
import secrets
import threading
import time
import uuid
import calendar
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

import msal
import pyodbc
import requests
from dotenv import load_dotenv
from flask import (Flask, Response, abort, flash, jsonify, redirect, render_template, request, send_file,
                   session, url_for)
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer

load_dotenv()

# --------------------------------------------------------------------------
# Configuration (all from environment / App Settings)
# --------------------------------------------------------------------------
TENANT_ID = os.environ.get("AZURE_TENANT_ID", "")
CLIENT_ID = os.environ.get("AZURE_CLIENT_ID", "")
CLIENT_SECRET = os.environ.get("AZURE_CLIENT_SECRET", "")

# Mailbox the app sends from (the app registration needs Graph Mail.Send for it)
SENDER_MAILBOX = os.environ.get("SENDER_MAILBOX", "")

# Public base URL of the portal INCLUDING the sub-path, no trailing slash. Every email button is built from it,
# e.g. http://mingalaba.hungerpangs.com:9000/dashboard/travel
APP_BASE_URL = os.environ.get("APP_BASE_URL", "").rstrip("/")

# Sub-path the portal is served under, e.g. /dashboard/travel (empty = site root). All routes, static files,
# redirects and url_for() links get this prefix.
_prefix = os.environ.get("APP_PREFIX", "").strip().strip("/")
APP_PREFIX = f"/{_prefix}" if _prefix else ""

# Reverse proxy in front of the portal (e.g. 127.0.0.1 when it runs on this server). Its X-Forwarded-For /
# X-Forwarded-Proto headers are then trusted, so the portal sees each visitor's real IP (used by User Management).
TRUSTED_PROXY = os.environ.get("TRUSTED_PROXY", "").strip()

# Waitress bind address. 0.0.0.0 = reachable from the internal network on every interface of this server
HOST = os.environ.get("HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT", "8080"))

# Signs the Approve/Reject links in manager emails. Must stay constant, or links already sent stop working.
SECRET_KEY = os.environ.get("SECRET_KEY", "")
REVIEW_LINK_MAX_AGE = int(os.environ.get("REVIEW_LINK_MAX_AGE_DAYS", "30")) * 86400

# Admin Team (Jigna, Mittal, Sagar) and CC (Pooja) - comma separated email addresses. (The variable names
# keep the old "Travel Desk" name so existing .env files keep working.)
TRAVEL_DESK_TO = [e.strip() for e in os.environ.get("TRAVEL_DESK_TO", "").split(",") if e.strip()]
TRAVEL_DESK_CC = [e.strip() for e in os.environ.get("TRAVEL_DESK_CC", "").split(",") if e.strip()]

# SQL Server
DB_DRIVER = os.environ.get("DB_DRIVER", "ODBC Driver 18 for SQL Server")
DB_SERVER = os.environ.get("DB_SERVER", "")
DB_NAME = os.environ.get("DB_NAME", "")
DB_USER = os.environ.get("DB_USER", "")
DB_PASSWORD = os.environ.get("DB_PASSWORD", "")
DB_PORT = os.environ.get("DB_PORT", "1433")
DB_TRUST_CERT = os.environ.get("DB_TRUST_SERVER_CERT", "yes")  # "yes" for self-signed certs
# Optional: a full ODBC connection string overrides the individual DB_* values above
DB_CONNECTION_STRING = os.environ.get("DB_CONNECTION_STRING", "")

# The portal's own employee -> manager routing table (migration 005), maintained by HR in User Management.
# dbo.checklist_empmanager_mapping belongs to the checklist apps: it was only read once, to seed this table.
MAPPING_TABLE = os.environ.get("MAPPING_TABLE", "dbo.TravelApproverMapping")
# Email columns inside the mapping table itself
EMP_EMAIL_COL = os.environ.get("EMP_EMAIL_COL", "EmpEmail")
MGR_EMAIL_COL = os.environ.get("MGR_EMAIL_COL", "ManagerEmail")

# HR master (name, email, status, Department/Designation ids) and the id -> name lookup tables
EMPLOYEE_TABLE = os.environ.get("EMPLOYEE_TABLE", "dbo.tblEmployeeMaster")
DEPARTMENT_TABLE = os.environ.get("DEPARTMENT_TABLE", "dbo.checklists_department")
DESIGNATION_TABLE = os.environ.get("DESIGNATION_TABLE", "dbo.checklists_designation")

# Optional login gate for the form (HTTP Basic). Leave PORTAL_PASSWORD empty to disable (local dev only).
PORTAL_USER = os.environ.get("PORTAL_USER", "")
PORTAL_PASSWORD = os.environ.get("PORTAL_PASSWORD", "")

# Optional image for the header. Empty (default) shows the built-in Hunger Pangs mark.
LOGO_URL = os.environ.get("LOGO_URL", "")

# User Management (HR edits the employee -> manager mapping). Empty password = section disabled.
HR_ADMIN_PASSWORD = os.environ.get("HR_ADMIN_PASSWORD", "")
# Who may open /user-management, as comma-separated networks. Default: office/private networks only, because
# the portal itself is reachable from the internet. Set to "any" to allow everyone (HTTPS strongly advised).
HR_ALLOWED_NETWORKS = os.environ.get(
    "HR_ALLOWED_NETWORKS", "10.0.0.0/8,172.16.0.0/12,192.168.0.0/16,127.0.0.0/8,::1/128")
HR_SESSION_MINUTES = int(os.environ.get("HR_SESSION_MINUTES", "30"))

ROOM_CATEGORIES = {"Single Occupancy", "Double Occupancy"}
REFUND_STATUSES = ("Refundable", "Non-refundable", "Partial refundable")

# Workflow: employee submits -> Admin Team proposes choices -> manager approves one/rejects
#           -> Admin Team makes the final booking and uploads the voucher -> employee gets it (manager in CC)
STATUS_BOOKING = "Pending Booking"            # waiting for the Admin Team to propose the booking
STATUS_MANAGER = "Pending Manager Approval"   # choices proposed, waiting for the manager
STATUS_FINAL = "Pending Final Booking"        # approved, waiting for the Admin Team to book and upload the voucher
STATUS_BOOKED = "Booked"                      # voucher emailed to the employee
DIRECT_NOTE = "No reporting manager on record: approved when the Admin Team confirmed the booking"
MAX_BOOKING_OPTIONS = 5  # choices the Admin Team can offer the manager
MAX_VOUCHER_BYTES = 3 * 1024 * 1024  # Graph sendMail takes attachments up to 3 MB in a single request
VOUCHER_TABLE = "dbo.TravelBookingVouchers"
GRAPH_SCOPE = ["https://graph.microsoft.com/.default"]

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_VOUCHER_BYTES + 512 * 1024  # voucher upload plus form fields
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("travel-portal")

if not SECRET_KEY:
    SECRET_KEY = secrets.token_urlsafe(32)
    log.warning("SECRET_KEY is not set: approval links will stop working when the server restarts")
# Separate salts: an Admin Team link can never be used as a manager approval link, and vice versa
_review_signer = URLSafeTimedSerializer(SECRET_KEY, salt="travel-review")
_desk_signer = URLSafeTimedSerializer(SECRET_KEY, salt="travel-desk")
_voucher_signer = URLSafeTimedSerializer(SECRET_KEY, salt="travel-voucher")

# Signed session cookie for the User Management unlock (HttpOnly; not readable by page scripts)
app.secret_key = SECRET_KEY
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_NAME="tp_session",
                  # Scope the cookie to the portal's sub-path so it never mixes with other apps on the same host
                  SESSION_COOKIE_PATH=APP_PREFIX or "/")


class PrefixMiddleware:
    """Serve the app under APP_PREFIX: /dashboard/travel/review/... reaches the /review/... route, and
    url_for() builds /dashboard/travel/... links (via SCRIPT_NAME). Anything outside the prefix is 404."""

    def __init__(self, wsgi_app, prefix):
        self.wsgi_app, self.prefix = wsgi_app, prefix

    def __call__(self, environ, start_response):
        path = environ.get("PATH_INFO", "")
        if path == self.prefix:  # /dashboard/travel -> /dashboard/travel/
            query = environ.get("QUERY_STRING")
            start_response("308 Permanent Redirect",
                           [("Location", self.prefix + "/" + ("?" + query if query else "")), ("Content-Length", "0")])
            return [b""]
        if path.startswith(self.prefix + "/"):
            environ["SCRIPT_NAME"] = environ.get("SCRIPT_NAME", "") + self.prefix
            environ["PATH_INFO"] = path[len(self.prefix):]
            return self.wsgi_app(environ, start_response)
        start_response("404 Not Found", [("Content-Type", "text/plain"), ("Content-Length", "9")])
        return [b"Not Found"]


if APP_PREFIX:
    app.wsgi_app = PrefixMiddleware(app.wsgi_app, APP_PREFIX)
    if APP_BASE_URL and not APP_BASE_URL.endswith(APP_PREFIX):
        log.warning("APP_BASE_URL (%s) should end with APP_PREFIX (%s) or email links will miss the sub-path",
                    APP_BASE_URL, APP_PREFIX)


# --------------------------------------------------------------------------
# Database helpers
# --------------------------------------------------------------------------
_IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*(\.[A-Za-z_][A-Za-z0-9_]*)?$")


def _q(identifier):
    """Validate and bracket-quote a table/column name that comes from configuration."""
    if not _IDENT.match(identifier):
        raise RuntimeError(f"Invalid SQL identifier in configuration: {identifier!r}")
    return ".".join(f"[{part}]" for part in identifier.split("."))


def get_conn():
    if DB_CONNECTION_STRING:
        return pyodbc.connect(DB_CONNECTION_STRING, timeout=15)
    conn_str = (
        f"DRIVER={{{DB_DRIVER}}};SERVER={DB_SERVER},{DB_PORT};DATABASE={DB_NAME};"
        f"UID={DB_USER};PWD={DB_PASSWORD};"
        f"Encrypt=yes;TrustServerCertificate={DB_TRUST_CERT};Connection Timeout=15;"
    )
    return pyodbc.connect(conn_str)


def _clean(v):
    return str(v).strip() if v is not None and str(v).strip() else None


def fetch_employee(emp_code):
    """
    Employee details from the HR master plus their reporting manager from the mapping table.
    Emails come from the mapping table, falling back to the HR master.
    Returns None if the code is in neither table. NeedsApproval is False when the employee
    has no reporting manager on record (HODs / senior staff).
    """
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT e.EmpCode, e.EmployeeName, e.EmployeeEmail, e.Status, dep.name, des.name,
                   m.ManagerCode, m.EmpEmail, m.ManagerEmail,
                   mgr.EmployeeName, mgr.EmployeeEmail, mdes.name
            FROM (SELECT CAST(? AS nvarchar(30)) AS EmpCode) k
            LEFT JOIN {_q(EMPLOYEE_TABLE)} e ON e.EmpCode = k.EmpCode
            LEFT JOIN {_q(DEPARTMENT_TABLE)} dep ON CAST(dep.id AS varchar(20)) = e.Department
            LEFT JOIN {_q(DESIGNATION_TABLE)} des ON CAST(des.id AS varchar(20)) = e.Designation
            OUTER APPLY (
                SELECT TOP 1 ManagerCode, {_q(EMP_EMAIL_COL)} AS EmpEmail, {_q(MGR_EMAIL_COL)} AS ManagerEmail
                FROM {_q(MAPPING_TABLE)} WHERE EmpCode = k.EmpCode ORDER BY id DESC
            ) m
            LEFT JOIN {_q(EMPLOYEE_TABLE)} mgr ON mgr.EmpCode = m.ManagerCode
            LEFT JOIN {_q(DESIGNATION_TABLE)} mdes ON CAST(mdes.id AS varchar(20)) = mgr.Designation
            """,
            emp_code,
        )
        row = cur.fetchone()
    (master_code, name, email, status, department, designation,
     mgr_code, map_emp_email, map_mgr_email, mgr_name, mgr_email, mgr_designation) = row
    mgr_code = _clean(mgr_code)
    if not master_code and not mgr_code:
        return None
    return {
        "EmpCode": _clean(master_code) or emp_code,  # HR casing, e.g. "h189" -> "H189"
        "EmpName": _clean(name),
        "EmpEmail": _clean(map_emp_email) or _clean(email),
        "Active": not master_code or _clean(status) == "Active",
        "Designation": _clean(designation),
        "Department": _clean(department),
        "NeedsApproval": bool(mgr_code),
        "ManagerCode": mgr_code,
        "ManagerName": _clean(mgr_name) or mgr_code,
        "ManagerDesignation": _clean(mgr_designation),
        "ManagerEmail": _clean(map_mgr_email) or _clean(mgr_email),
    }


def employee_problem(emp):
    """A reason this employee cannot submit yet (shown at lookup time), or None."""
    if not emp["Active"]:
        return "This employee code is not active in HR records."
    if emp["NeedsApproval"] and not emp["ManagerEmail"]:
        return (f"Your manager ({emp['ManagerName']}) has no email address on record, so the approval "
                "email cannot be sent. Please ask HR to update it.")
    if not emp["EmpEmail"]:
        return "Your email address is not on record. Please ask HR to update it."
    return None


def insert_request(req):
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO TravelRequests
              (RequestId, EmpCode, EmpName, EmpEmail, ManagerCode, ManagerName, ManagerEmail,
               Designation, Department, FromDate, ToDate, PlaceFrom, PlaceTo,
               RoomCategory, PreferredLocation, PreferredHotel, ReasonOfTravel, Status)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            req["RequestId"], req["EmpCode"], req["EmpName"], req["EmpEmail"],
            req["ManagerCode"], req["ManagerName"], req["ManagerEmail"],
            req["Designation"], req["Department"], req["FromDate"], req["ToDate"],
            req["PlaceFrom"], req["PlaceTo"], req["RoomCategory"], req["PreferredLocation"],
            req["PreferredHotel"] or None, req["ReasonOfTravel"], STATUS_BOOKING,
        )
        conn.commit()


def set_status(request_id, status):
    with get_conn() as conn:
        conn.cursor().execute(
            "UPDATE TravelRequests SET Status = ? WHERE RequestId = ?", status, request_id
        )
        conn.commit()


def load_request(request_id):
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute("SELECT * FROM TravelRequests WHERE RequestId = ?", request_id)
        row = cur.fetchone()
        if not row:
            return None
        cols = [c[0] for c in cur.description]
        return dict(zip(cols, row))


def advance_status(request_id, from_status, to_status):
    """Atomically move a request from one status to the next. Returns True if this call won."""
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute("UPDATE TravelRequests SET Status = ? WHERE RequestId = ? AND Status = ?",
                    to_status, request_id, from_status)
        won = cur.rowcount == 1
        conn.commit()
        return won


def save_booking(request_id, options, to_status, comment=None):
    """Store the Admin Team's booking choices (JSON) and move Pending Booking -> to_status atomically.
    to_status is Pending Manager Approval, or Pending Final Booking when there is no reporting manager - then the
    single choice is also written to BookingArea/PropertyName/Cost as the approved booking.
    Returns True if this call won."""
    approved = to_status == STATUS_FINAL
    sets, args = ["BookingOptions = ?", "Status = ?"], [options_json(options), to_status]
    if approved:
        final = options[0]
        sets += ["BookingArea = ?", "PropertyName = ?", "Cost = ?", "SelectedOption = 1",
                 "ActionedAt = SYSUTCDATETIME()", "ActionComment = ?"]
        args += [final["area"], final["property"], final["cost"], comment]
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"UPDATE TravelRequests SET {', '.join(sets)} WHERE RequestId = ? AND Status = ?",
                    *args, request_id, STATUS_BOOKING)
        won = cur.rowcount == 1
        conn.commit()
        return won


def claim_decision(request_id, status, comment, from_status=None, option=None, option_no=None):
    """Atomically move a request awaiting a decision to Pending Final Booking (approved) or Rejected. On approval,
    the manager's chosen booking choice (1-based option_no) becomes BookingArea/PropertyName/Cost in the same update.
    Returns True if this call won."""
    chosen = ", BookingArea = ?, PropertyName = ?, Cost = ?, SelectedOption = ?" if option else ""
    chosen_args = [option["area"], option["property"], option["cost"], option_no] if option else []
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            "UPDATE TravelRequests SET Status = ?, ActionedAt = SYSUTCDATETIME(), ActionComment = ?, "
            f"RejectionReason = ?{chosen} WHERE RequestId = ? AND Status = ?",
            status, comment, (comment if status == "Rejected" else None), *chosen_args, request_id,
            from_status or STATUS_MANAGER,
        )
        won = cur.rowcount == 1
        conn.commit()
        return won


def save_voucher(request_id, file_name, content):
    """Store the final booking voucher and move Pending Final Booking -> Booked in one transaction.
    Returns True if this call won. A re-upload (after a failed email) replaces the earlier file."""
    with get_conn() as conn:
        conn.autocommit = False
        cur = conn.cursor()
        cur.execute("UPDATE TravelRequests SET Status = ? WHERE RequestId = ? AND Status = ?",
                    STATUS_BOOKED, request_id, STATUS_FINAL)
        if cur.rowcount != 1:
            conn.rollback()
            return False
        cur.execute(f"DELETE FROM {_q(VOUCHER_TABLE)} WHERE RequestId = ?", request_id)
        cur.execute(f"INSERT INTO {_q(VOUCHER_TABLE)} (RequestId, FileName, Content) VALUES (?, ?, ?)",
                    request_id, file_name, pyodbc.Binary(content))
        conn.commit()
        return True


def load_voucher(request_id):
    """(file name, PDF bytes, uploaded at UTC) or None."""
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT FileName, Content, UploadedAt FROM {_q(VOUCHER_TABLE)} WHERE RequestId = ?", request_id)
        row = cur.fetchone()
    return (row[0], bytes(row[1]), row[2]) if row else None


def export_requests(start, end, basis):
    """Requests whose request date (basis 'created', India time) or check-in date (basis 'checkin') falls in
    [start, end] inclusive, oldest first, with the voucher upload time. Never reads the voucher files."""
    column = "CAST(DATEADD(minute, 330, r.CreatedAt) AS date)" if basis == "created" else "r.FromDate"
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT r.*, v.UploadedAt AS VoucherUploadedAt, v.FileName AS VoucherFileName
            FROM TravelRequests r
            LEFT JOIN {_q(VOUCHER_TABLE)} v ON v.RequestId = r.RequestId
            WHERE {column} BETWEEN ? AND ?
            ORDER BY r.BookingRefID
            """,
            start, end,
        )
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, row)) for row in cur.fetchall()]


# --------------------------------------------------------------------------
# Microsoft Graph helpers
# --------------------------------------------------------------------------
_msal_app = None


def get_graph_token():
    global _msal_app
    if _msal_app is None:
        _msal_app = msal.ConfidentialClientApplication(
            CLIENT_ID,
            authority=f"https://login.microsoftonline.com/{TENANT_ID}",
            client_credential=CLIENT_SECRET,
        )
    result = _msal_app.acquire_token_for_client(scopes=GRAPH_SCOPE)  # MSAL caches tokens
    if "access_token" not in result:
        raise RuntimeError(f"Graph token error: {result.get('error_description', result)}")
    return result["access_token"]


def send_mail(to, subject, html_body, cc=None, attachments=None):
    """attachments: [(file name, content type, bytes)], 3 MB in total at most."""
    message = {
        "subject": subject,
        "body": {"contentType": "HTML", "content": html_body},
        "toRecipients": [{"emailAddress": {"address": a}} for a in to],
        "ccRecipients": [{"emailAddress": {"address": a}} for a in (cc or [])],
    }
    if attachments:
        message["attachments"] = [
            {"@odata.type": "#microsoft.graph.fileAttachment", "name": name, "contentType": ctype,
             "contentBytes": base64.b64encode(data).decode("ascii")}
            for name, ctype, data in attachments
        ]
    payload = {"message": message, "saveToSentItems": True}
    resp = requests.post(
        f"https://graph.microsoft.com/v1.0/users/{SENDER_MAILBOX}/sendMail",
        headers={"Authorization": f"Bearer {get_graph_token()}", "Content-Type": "application/json"},
        json=payload,
        timeout=30,
    )
    if resp.status_code != 202:
        raise RuntimeError(f"sendMail failed {resp.status_code}: {resp.text}")


# --------------------------------------------------------------------------
# Presentation helpers
# --------------------------------------------------------------------------
def fmt_date(d):
    if isinstance(d, (date, datetime)):
        return d.strftime("%d-%b-%Y")
    return str(d)


def booking_ref(r):
    """Short reference for emails and pages, e.g. '#10001'. Falls back to the GUID prefix before migration 003."""
    if r.get("BookingRefID"):
        return f"#{r['BookingRefID']}"
    return "#" + str(r["RequestId"])[:8].upper()


def with_ref(r, subject):
    return f"Booking Ref: {booking_ref(r)} - {subject}"


# Email building blocks. Inline styles only (email clients ignore <style> blocks); tables are 100% wide up to
# 600px so they fill a phone screen and stay readable on desktop.
_TABLE = ("role='presentation' width='100%' cellpadding='0' cellspacing='0' "
          "style='width:100%;max-width:600px;border-collapse:collapse;font-family:Segoe UI,Arial,sans-serif'")
_BOX = ("role='presentation' width='100%' cellpadding='0' cellspacing='0' "
        "style='width:100%;max-width:600px;border-collapse:separate;margin:0 0 18px;border-radius:12px;{}'")


def ref_badge_html(r):
    return ("<p style='margin:0 0 16px'><span style='display:inline-block;padding:5px 12px;border-radius:999px;"
            "background:#f7efdd;color:#681313;font-weight:700;font-size:13px'>"
            f"Booking Ref: {html.escape(booking_ref(r))}</span></p>")


def fmt_duration(from_date, to_date):
    """'3 nights, 4 days' - same wording as the form's Trip summary."""
    nights = (to_date - from_date).days
    days = nights + 1
    return f"{nights} night{'' if nights == 1 else 's'}, {days} day{'' if days == 1 else 's'}"


def details_table_html(r, booking=True):
    """Employee and trip details. booking=False leaves out the Admin Team's booking (shown in its own box)."""
    rows = [
        ("Booking Ref", booking_ref(r)),
        ("Emp Code", r["EmpCode"]),
        ("Name", r["EmpName"]),
        ("Designation", r["Designation"]),
        ("Department", r["Department"]),
        ("From (Date)", fmt_date(r["FromDate"])),
        ("To (Date)", fmt_date(r["ToDate"])),
        ("Duration", fmt_duration(r["FromDate"], r["ToDate"])),
        ("Travelling From", r["PlaceFrom"]),
        ("Travelling To", r["PlaceTo"]),
        ("Reason of Travel", r.get("ReasonOfTravel") or "-"),
        ("Category of Room", r["RoomCategory"]),
        ("Preferred Location", r["PreferredLocation"] or "-"),
        ("Preferred Hotel", r.get("PreferredHotel") or "-"),
    ]
    if booking and r.get("PropertyName"):
        rows += booking_rows(r)
    last = len(rows) - 1
    body = "".join(
        "<tr>"
        f"<td width='40%' style='width:40%;padding:10px 14px;background:#f6efe6;color:#6b5a4e;font-size:13px;"
        f"font-weight:600;vertical-align:top;{'' if i == last else 'border-bottom:1px solid #ece2d6;'}'>"
        f"{html.escape(k)}</td>"
        f"<td style='padding:10px 14px;color:#3b2f2a;font-size:14px;font-weight:600;vertical-align:top;"
        f"word-break:break-word;{'' if i == last else 'border-bottom:1px solid #ece2d6;'}'>{html.escape(str(v))}</td>"
        "</tr>"
        for i, (k, v) in enumerate(rows)
    )
    # Rounded warm card; overflow:hidden keeps the tinted label column inside the rounded corners
    return (
        "<table role='presentation' width='100%' cellpadding='0' cellspacing='0' "
        "style='width:100%;max-width:600px;border-collapse:separate;border-spacing:0;background:#fdfaf6;"
        "border:1px solid #e6ddd3;border-radius:12px;overflow:hidden;font-family:Segoe UI,Arial,sans-serif'>"
        f"{body}</table>"
    )


def booking_box_html(r, title="Booking arranged by the Admin Team", rows=None, emphasis=False):
    """A booking (Place/Area, Hotel, Total Cost) in a soft, brand-tinted box.
    rows defaults to the approved booking; emphasis=True gives the approved choice a stronger border."""
    rows = "".join(
        "<tr>"
        f"<td width='40%' style='width:40%;padding:6px 12px 6px 0;color:#8a5a4c;font-size:13px'>{html.escape(k)}</td>"
        "<td style='padding:6px 0;color:#1d1d1d;font-size:15px;font-weight:700;word-break:break-word'>"
        f"{html.escape(v)}</td></tr>"
        for k, v in (rows if rows is not None else booking_rows(r))
    )
    border = "border:2px solid #c54b36" if emphasis else "border:1px solid #f3dccf"
    return (
        f"<table {_BOX.format('background:#fdf3ec;' + border)}>"
        "<tr><td style='padding:14px 18px;font-family:Segoe UI,Arial,sans-serif'>"
        "<div style='font-size:12px;font-weight:700;letter-spacing:.06em;text-transform:uppercase;color:#c54b36;"
        f"margin-bottom:6px'>{html.escape(title)}</div>"
        f"<table {_TABLE}>{rows}</table>"
        "</td></tr></table>"
    )


def choice_label(n, total):
    """Just the number ('2') when the Admin Team offered several; no number at all for a single choice."""
    return str(n) if total > 1 else "Proposed booking"


def options_html(r, heading_note="", admin=False):
    """Every choice the Admin Team offered, one box each (numbered 1, 2, ... when there are several). admin=True adds the
    Admin-only fields; only emails to the Admin Team pass it."""
    opts = booking_options(r)
    return "".join(booking_box_html(r, choice_label(i, len(opts)) + heading_note, option_rows(o, admin))
                   for i, o in enumerate(opts, 1))


def approved_box_html(r, title, admin=False):
    """The choice the manager approved, highlighted."""
    n, total = selected_option_no(r), len(booking_options(r))
    label = f"{title} ({n})" if n and total > 1 else title
    return booking_box_html(r, label, booking_rows(r, admin), emphasis=True)


def reason_box_html(label, text):
    """The manager's reason/comment, highlighted."""
    return (
        f"<table {_BOX.format('background:#fbeae6;border:1px solid #f0c4b9')}>"
        "<tr><td style='padding:14px 18px;font-family:Segoe UI,Arial,sans-serif'>"
        "<div style='font-size:12px;font-weight:700;letter-spacing:.06em;text-transform:uppercase;color:#9b3523;"
        f"margin-bottom:4px'>{html.escape(label)}</div>"
        f"<div style='font-size:15px;color:#1d1d1d;word-break:break-word'>{html.escape(text)}</div>"
        "</td></tr></table>"
    )


def email_html(*parts):
    """Wrap an email body: 600px max, full width on phones."""
    return f"<div style='{EMAIL_FONT};max-width:600px;line-height:1.5'>" + "".join(parts) + "</div>"


def fmt_inr(amount):
    """Indian digit grouping: 123456.5 -> '₹1,23,456.50'."""
    whole, frac = f"{Decimal(amount):.2f}".split(".")
    head, tail = whole[:-3], whole[-3:]
    while len(head) > 2:
        tail = head[-2:] + "," + tail
        head = head[:-2]
    return "₹" + (head + "," + tail if head else tail) + "." + frac


COST_LABEL = "Total Cost of stay with GST (approx)"
# Filled in by the Admin Team for each choice; shown only to the Admin Team (never to the manager or employee)
ADMIN_FIELDS = (("occupancy", "Type of occupancy"), ("refund", "Refund status"), ("remarks", "Remarks"))


def booking_rows(r, admin=False):
    """The approved booking as (label, value) pairs. admin=True adds the Admin-only fields of the chosen choice."""
    rows = [
        ("Place/Area", r.get("BookingArea") or "-"),
        ("Hotel", r.get("PropertyName") or "-"),
        (COST_LABEL, fmt_inr(r["Cost"]) if r.get("Cost") is not None else "-"),
    ]
    if admin:
        rows += [(label, chosen_option(r).get(key) or "-") for key, label in ADMIN_FIELDS]
    return rows


def options_json(options):
    return json.dumps([{"area": o["area"], "property": o["property"], "cost": f"{o['cost']:.2f}",
                        **{key: o.get(key) or "" for key, _ in ADMIN_FIELDS}} for o in options],
                      ensure_ascii=False)


def booking_options(r):
    """All choices the Admin Team offered, as [{'area', 'property', 'cost': Decimal, 'occupancy', 'refund',
    'remarks'}]. Requests saved before multiple choices existed have just their single booking."""
    raw = r.get("BookingOptions")
    if raw:
        try:
            return [{"area": o["area"], "property": o["property"], "cost": Decimal(o["cost"]),
                     **{key: o.get(key) or "" for key, _ in ADMIN_FIELDS}} for o in json.loads(raw)]
        except (ValueError, KeyError, TypeError, InvalidOperation):
            log.warning("Unreadable BookingOptions for %s", r.get("RequestId"))
    if r.get("PropertyName"):
        return [{"area": r.get("BookingArea") or "", "property": r["PropertyName"], "cost": r.get("Cost") or Decimal(0),
                 **{key: "" for key, _ in ADMIN_FIELDS}}]
    return []


def option_rows(o, admin=False):
    rows = [("Place/Area", o["area"] or "-"), ("Hotel", o["property"] or "-"), (COST_LABEL, fmt_inr(o["cost"]))]
    if admin:
        rows += [(label, o.get(key) or "-") for key, label in ADMIN_FIELDS]
    return rows


def selected_option_no(r):
    """1-based number of the choice the manager approved (1 for single-choice and older requests)."""
    return r.get("SelectedOption") or (1 if r.get("PropertyName") else None)


def chosen_option(r):
    """The approved choice with its Admin-only fields, or {}."""
    n, opts = selected_option_no(r), booking_options(r)
    return opts[n - 1] if n and n <= len(opts) else {}


def is_approved(r):
    """Approved by the manager (or no manager needed): waiting for the voucher, booked, or approved before the
    voucher step existed."""
    return r.get("Status") in (STATUS_FINAL, STATUS_BOOKED, "Approved")


app.jinja_env.globals.update(booking_ref=booking_ref, fmt_duration=fmt_duration, booking_options=booking_options,
                             option_rows=option_rows, selected_option_no=selected_option_no, choice_label=choice_label,
                             is_approved=is_approved, refund_statuses=REFUND_STATUSES, cost_label=COST_LABEL,
                             max_booking_options=MAX_BOOKING_OPTIONS,
                             max_voucher_mb=MAX_VOUCHER_BYTES // (1024 * 1024))  # same values as the emails


# --------------------------------------------------------------------------
# Signed action links and workflow emails - plain HTML, no Actionable Messages
# --------------------------------------------------------------------------
def make_token(signer, request_id):
    return signer.dumps(str(request_id))


def check_token(signer, request_id, token):
    """True if the token was issued by this signer for this request and has not expired."""
    try:
        return signer.loads(token, max_age=REVIEW_LINK_MAX_AGE) == str(request_id)
    except (SignatureExpired, BadSignature):
        return False


def desk_url(request_id):
    return f"{APP_BASE_URL}/travel-desk-confirm/{request_id}?t={make_token(_desk_signer, request_id)}"


def review_url(request_id, decision, quick=False):
    """quick=True: the Approve button for a single choice - opens the minimal confirm page."""
    url = f"{APP_BASE_URL}/review/{request_id}?t={make_token(_review_signer, request_id)}&decision={decision}"
    return url + "&quick=1" if quick else url


def voucher_url(request_id):
    return f"{APP_BASE_URL}/final-booking/{request_id}?t={make_token(_voucher_signer, request_id)}"


EMAIL_FONT = "font-family:Segoe UI,Arial,sans-serif;font-size:14px;color:#1d1d1d"
LINK_NOTE = ("<p style='color:#7a6f66;font-size:12px'>The button opens the Travel Accommodation Portal "
             "in your browser.</p>")


def _email_button(label, url, color):
    return (
        f"<a href='{html.escape(url)}' style='display:inline-block;padding:11px 24px;margin:0 10px 6px 0;"
        f"background:{color};color:#ffffff;text-decoration:none;font-weight:600;border-radius:8px;"
        f"font-family:Segoe UI,Arial,sans-serif;font-size:14px'>{html.escape(label)}</a>"
    )


def desk_email_html(r):
    """Step 1 - Admin Team: find the hotel, then send the choices to the manager."""
    direct = not r.get("ManagerCode")
    then = ("Once arranged, click <b>Enter Booking Details</b> to add the booking. This employee has "
            "<b>no reporting manager on record</b>, so no approval is needed: you can make the final booking "
            "straight after."
            if direct else
            "Once arranged, click <b>Enter Booking Details</b>.")
    return email_html(
        ref_badge_html(r),
        "<p>Dear Admin Team,</p>",
        f"<p>{html.escape(r['EmpName'])} ({html.escape(r['EmpCode'])}) has submitted a hotel booking request. "
        f"Please arrange the hotel. {then}</p>",
        details_table_html(r),
        "<p style='margin:22px 0 8px'>", _email_button("Enter Booking Details", desk_url(r["RequestId"]), "#c54b36"),
        "</p>",
        LINK_NOTE,
    )


def manager_email_html(r):
    """Step 2 - Manager: approve the Admin Team's choice (or pick one of several), or reject."""
    rid = r["RequestId"]
    n = len(booking_options(r))
    if n > 1:
        ask = (f"The Admin Team has found <b>{n} hotel choices</b>. Please choose one to approve, "
               "or reject the request. A reason is required to reject.")
        approve = _email_button("Choose & Approve", review_url(rid, "approve"), "#7f8e28")
    else:
        ask = "Please approve or reject the booking below. A reason is required to reject."
        approve = _email_button("Approve", review_url(rid, "approve", quick=True), "#7f8e28")
    return email_html(
        ref_badge_html(r),
        f"<p>Dear {html.escape(r['ManagerName'])},</p>",
        f"<p>The Admin Team has arranged a hotel for {html.escape(r['EmpName'])} "
        f"({html.escape(r['EmpCode'])}). {ask}</p>",
        options_html(r),
        details_table_html(r, booking=False),
        "<p style='margin:22px 0 8px'>", approve,
        _email_button("Reject", review_url(rid, "reject"), "#c54b36"),
        "</p>",
        LINK_NOTE,
    )


def final_booking_email_html(r):
    """Step 3 - Admin Team: the booking is approved; book it and upload the voucher."""
    direct = not r.get("ManagerCode")
    multi = len(booking_options(r)) > 1
    how = ("no manager approval was required (no reporting manager on record)" if direct
           else f"approved by {html.escape(r['ManagerName'])}")
    chose = (f" The manager chose <b>{selected_option_no(r)}</b>; please release the others."
             if multi else "")
    return email_html(
        ref_badge_html(r),
        "<p>Dear Admin Team,</p>",
        f"<p>The hotel booking for {html.escape(r['EmpName'])} ({html.escape(r['EmpCode'])}) is <b>approved</b>: "
        f"{how}.{chose}</p>",
        "<p>Please make the final booking, then click <b>Upload Booking Voucher</b> to email the voucher (PDF).</p>",
        approved_box_html(r, "Approved booking", admin=True),
        details_table_html(r, booking=False),
        "<p style='margin:22px 0 8px'>",
        _email_button("Upload Booking Voucher", voucher_url(r["RequestId"]), "#c54b36"), "</p>",
        LINK_NOTE,
    )


def voucher_email_html(r):
    """Step 4 - Employee (manager in CC): the booking is confirmed; the voucher is attached."""
    dates = f"{fmt_date(r['FromDate'])} to {fmt_date(r['ToDate'])}"
    return email_html(
        ref_badge_html(r),
        f"<p>Dear {html.escape(r['EmpName'])},</p>",
        f"<p>Your hotel booking in {html.escape(r['PlaceTo'])} ({dates}) is <b>confirmed</b>. "
        "The booking voucher is attached to this email (PDF).</p>",
        approved_box_html(r, "Your booking"),
        details_table_html(r, booking=False),
    )


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------
# Admin Team and manager links carry their own signed token; /healthz is for monitoring; /static/ is the logo
_OPEN_PREFIXES = ("/travel-desk-confirm/", "/desk/", "/review/", "/final-booking/", "/healthz", "/static/")


@app.before_request
def require_portal_login():
    if not PORTAL_PASSWORD or request.path.startswith(_OPEN_PREFIXES):
        return None
    auth = request.authorization
    if (auth
            and secrets.compare_digest((auth.username or "").encode(), PORTAL_USER.encode())
            and secrets.compare_digest((auth.password or "").encode(), PORTAL_PASSWORD.encode())):
        return None
    return Response("Login required", 401, {"WWW-Authenticate": 'Basic realm="Travel Portal"'})


@app.errorhandler(413)
def too_large(_):
    mb = MAX_VOUCHER_BYTES // (1024 * 1024)
    return Response(f"The file is too large. The booking voucher must be a PDF of at most {mb} MB. "
                    "Go back and choose a smaller file.", 413, {"Content-Type": "text/plain; charset=utf-8"})


@app.get("/")
def index():
    today = date.today().isoformat()
    return render_template("index.html", today=today, logo_url=LOGO_URL, show_hr_link=True)


@app.get("/healthz")
def healthz():
    return "ok", 200


@app.get("/api/employee/<emp_code>")
def api_employee(emp_code):
    if not re.fullmatch(r"[A-Za-z0-9\-_/.]{1,30}", emp_code):
        return jsonify({"found": False, "error": "Invalid employee code"}), 400
    try:
        emp = fetch_employee(emp_code.strip())
    except Exception:  # noqa: BLE001
        log.exception("Employee lookup failed")
        return jsonify({"found": False, "error": "Lookup failed"}), 500
    if not emp:
        return jsonify({"found": False}), 404
    return jsonify({
        "found": True,
        "empName": emp["EmpName"] or "",
        "designation": emp["Designation"] or "",
        "department": emp["Department"] or "",
        "needsApproval": emp["NeedsApproval"],
        "managerName": emp["ManagerName"] if emp["NeedsApproval"] else None,
        "managerDesignation": emp["ManagerDesignation"] if emp["NeedsApproval"] else None,
        "problem": employee_problem(emp),
    })


@app.post("/submit-travel")
def submit_travel():
    f = request.form if request.form else (request.get_json(silent=True) or {})
    get = lambda k: str(f.get(k, "")).strip()  # noqa: E731

    emp_code = get("emp_code")
    designation, department = get("designation"), get("department")
    place_from, place_to = get("place_from"), get("place_to")
    room, pref_loc, pref_hotel = get("room_category"), get("preferred_location"), get("preferred_hotel")
    travel_reason = get("travel_reason")

    missing = [n for n, v in {
        "Emp Code": emp_code, "Designation": designation, "Department": department,
        "From (Date)": get("from_date"), "To (Date)": get("to_date"),
        "Travelling From": place_from, "Travelling To": place_to, "Reason of Travel": travel_reason,
        "Category of Room": room,
    }.items() if not v]
    if missing:
        return jsonify({"ok": False, "error": "Missing: " + ", ".join(missing)}), 400

    try:
        from_date = datetime.strptime(get("from_date"), "%Y-%m-%d").date()
        to_date = datetime.strptime(get("to_date"), "%Y-%m-%d").date()
    except ValueError:
        return jsonify({"ok": False, "error": "Invalid date format"}), 400
    if to_date < from_date:
        return jsonify({"ok": False, "error": "'To' date cannot be before 'From' date"}), 400
    if room not in ROOM_CATEGORIES:
        return jsonify({"ok": False, "error": "Invalid room category"}), 400
    if len(travel_reason) > 500:
        return jsonify({"ok": False, "error": "Reason of Travel is too long (max 500 characters)"}), 400
    if max(len(designation), len(department), len(place_from), len(place_to), len(pref_loc)) > 150:
        return jsonify({"ok": False, "error": "A field is too long (max 150 characters)"}), 400
    if len(pref_hotel) > 200:
        return jsonify({"ok": False, "error": "Preferred Hotel name is too long (max 200 characters)"}), 400

    try:
        emp = fetch_employee(emp_code)
    except Exception:  # noqa: BLE001
        log.exception("DB error during employee fetch")
        return jsonify({"ok": False, "error": "Database error. Please try again."}), 500
    if not emp:
        return jsonify({"ok": False, "error": f"Employee code '{emp_code}' not found"}), 404
    problem = employee_problem(emp)
    if problem:
        return jsonify({"ok": False, "error": problem}), 422
    emp["EmpName"] = emp["EmpName"] or get("emp_name")
    if not emp["EmpName"]:
        return jsonify({"ok": False, "error": "Name of Person travelling is required"}), 400

    # HR-record values (name, designation, department, manager) win over what the browser sent
    req = {
        "RequestId": str(uuid.uuid4()),
        **{k: emp[k] for k in ("EmpCode", "EmpName", "EmpEmail", "ManagerCode")},
        "ManagerName": emp["ManagerName"] if emp["NeedsApproval"] else None,
        "ManagerEmail": emp["ManagerEmail"] if emp["NeedsApproval"] else "",  # column is NOT NULL
        "Designation": emp["Designation"] or designation, "Department": emp["Department"] or department,
        "FromDate": from_date, "ToDate": to_date,
        "PlaceFrom": place_from, "PlaceTo": place_to,
        "RoomCategory": room, "PreferredLocation": pref_loc, "PreferredHotel": pref_hotel,
        "ReasonOfTravel": travel_reason,
    }

    try:
        insert_request(req)
        req = load_request(req["RequestId"]) or req  # picks up the database-assigned BookingRefID
    except Exception:  # noqa: BLE001
        log.exception("DB error during insert")
        return jsonify({"ok": False, "error": "Could not save the request."}), 500

    try:
        send_mail(
            TRAVEL_DESK_TO,
            with_ref(req, f"Hotel booking request: {req['EmpName']} ({req['EmpCode']}) - "
                          f"{req['PlaceFrom']} to {req['PlaceTo']}"),
            desk_email_html(req),
            cc=TRAVEL_DESK_CC,
        )
    except Exception:  # noqa: BLE001
        log.exception("Failed to email the Admin Team")
        set_status(req["RequestId"], "MailFailed")
        return jsonify({"ok": False, "error": "Could not email the Admin Team. Please contact support."}), 502

    return jsonify({
        "ok": True,
        "requestId": str(req["RequestId"]),
        "bookingRef": booking_ref(req),
        "needsApproval": emp["NeedsApproval"],
    })


def notify_approval(r):
    """Approved (by the manager, or no manager needed): the Admin Team makes the final booking and uploads the
    voucher. The employee hears next when the voucher is emailed."""
    trip = f"{r['PlaceFrom']} to {r['PlaceTo']}"
    send_mail(TRAVEL_DESK_TO,
              with_ref(r, f"Approved - final booking needed: {r['EmpName']} ({r['EmpCode']}) - {trip}"),
              final_booking_email_html(r), cc=TRAVEL_DESK_CC)


def notify_rejection(r, comment):
    """Rejected: separate emails to the employee (with the manager's reason) and the Admin Team (cancel the
    tentative booking)."""
    who = f"{html.escape(r['EmpName'])} ({html.escape(r['EmpCode'])})"
    trip = f"{r['PlaceFrom']} to {r['PlaceTo']}"
    dates = f"{fmt_date(r['FromDate'])} to {fmt_date(r['ToDate'])}"
    emp_subject = with_ref(r, "Your hotel booking request was rejected")
    emp_body = email_html(
        ref_badge_html(r),
        f"<p>Dear {html.escape(r['EmpName'])},</p>",
        f"<p>Your hotel booking request from {html.escape(r['PlaceFrom'])} to {html.escape(r['PlaceTo'])} "
        f"({dates}) has been <b>rejected</b> by {html.escape(r['ManagerName'])}. "
        "The Admin Team will cancel the tentative booking.</p>",
        reason_box_html("Reason for rejection", comment or "-"),
        details_table_html(r, booking=False),
    )
    desk_subject = with_ref(r, f"Rejected - cancel booking: {r['EmpName']} ({r['EmpCode']}) - {trip}")
    desk_body = email_html(
        ref_badge_html(r),
        "<p>Dear Admin Team,</p>",
        f"<p>The hotel booking for {who} has been <b>rejected</b> by {html.escape(r['ManagerName'])}. "
        "<b>Please cancel the tentative booking below.</b></p>",
        reason_box_html("Reason for rejection", comment or "-"),
        options_html(r, " - cancel", admin=True),
        details_table_html(r, booking=False),
    )
    errors = []
    for to, cc, subject, body in (
        ([r["EmpEmail"]], None, emp_subject, emp_body),  # always straight to the employee
        (TRAVEL_DESK_TO, TRAVEL_DESK_CC, desk_subject, desk_body),
    ):
        try:
            send_mail(to, subject, body, cc=cc)
        except Exception as exc:  # noqa: BLE001 - still try the other recipient
            errors.append(exc)
    if errors:
        raise errors[0]


def _check_link(signer, request_id, template):
    """Validate the request id and signed token. Returns (token, error_response or None)."""
    token = request.values.get("t", "")
    try:
        uuid.UUID(request_id)
    except ValueError:
        abort(404)
    if not check_token(signer, request_id, token):
        return token, (render_template(template, state="invalid", logo_url=LOGO_URL), 403)
    return token, None


def _parse_cost(raw):
    try:
        cost = Decimal(raw.replace(",", "").replace("₹", "").strip()).quantize(Decimal("0.01"))
        return cost if Decimal("0") < cost < Decimal("10000000000") else None
    except (InvalidOperation, ValueError):
        return None


def parse_booking_options(form, allow_multiple=True):
    """Validate the Admin Team form. The first choice is required; every extra choice must be complete (Remarks are
    optional). Returns (rows for re-display, options, errors) where errors maps 'field-index' -> message."""
    fields = ("booking_area", "property_name", "cost", "occupancy", "refund", "remarks")
    lists = {k: form.getlist(k) for k in fields}
    n = max(max(len(v) for v in lists.values()), 1)
    val = lambda k, i, limit: (lists[k][i] if i < len(lists[k]) else "").strip()[:limit]  # noqa: E731
    rows, options, errors = [], [], {}
    for i in range(n):
        row = {"booking_area": val("booking_area", i, 150), "property_name": val("property_name", i, 200),
               "cost": val("cost", i, 20), "occupancy": val("occupancy", i, 100), "refund": val("refund", i, 30),
               "remarks": val("remarks", i, 300)}
        rows.append(row)
        cost = _parse_cost(row["cost"])
        if not row["booking_area"]:
            errors[f"booking_area-{i}"] = "Enter the place/area of the stay."
        if not row["property_name"]:
            errors[f"property_name-{i}"] = "Enter the hotel name."
        if cost is None:
            errors[f"cost-{i}"] = "Enter the total cost as a number, e.g. 12500."
        if not row["occupancy"]:
            errors[f"occupancy-{i}"] = "Enter the type of occupancy."
        if row["refund"] not in REFUND_STATUSES:
            errors[f"refund-{i}"] = "Choose the refund status."
        options.append({"area": row["booking_area"], "property": row["property_name"], "cost": cost,
                        "occupancy": row["occupancy"], "refund": row["refund"], "remarks": row["remarks"]})
    if n > (MAX_BOOKING_OPTIONS if allow_multiple else 1):
        errors["form"] = (f"You can offer at most {MAX_BOOKING_OPTIONS} choices." if allow_multiple else
                          "This employee has no reporting manager, so only one booking can be entered.")
    return rows, options, errors


def _voucher_link(request_id):
    """Relative link to the voucher page (for pages shown to the Admin Team)."""
    return url_for("final_booking", request_id=request_id, t=make_token(_voucher_signer, request_id))


@app.get("/desk/<request_id>")
def desk_legacy(request_id):
    """Links in emails sent before the booking-details form was added."""
    return redirect(url_for("travel_desk_confirm", request_id=request_id, t=request.args.get("t", "")), 301)


@app.route("/travel-desk-confirm/<request_id>", methods=["GET", "POST"])
def travel_desk_confirm(request_id):
    """Step 1 - the Admin Team enters the hotel choices and sends them to the manager.
    GET only shows the form (safe for mail link scanners); POST saves and emails the manager."""
    template = "travel_desk_confirm.html"
    token, err = _check_link(_desk_signer, request_id, template)
    if err:
        return err
    r = load_request(request_id)
    if not r:
        abort(404)
    saved = booking_options(r)  # e.g. after a failed manager email, show what was entered
    blank = {"booking_area": "", "property_name": "", "cost": "", "occupancy": "", "refund": "", "remarks": ""}
    page = dict(r=r, token=token, fmt_date=fmt_date, booking_rows=booking_rows, logo_url=LOGO_URL, errors={},
                voucher_link=None,
                rows=[{"booking_area": o["area"], "property_name": o["property"], "cost": f"{o['cost']:.2f}",
                       "occupancy": o["occupancy"], "refund": o["refund"], "remarks": o["remarks"]}
                      for o in saved] or [blank])

    if r["Status"] != STATUS_BOOKING:
        if r["Status"] == STATUS_FINAL:
            page["voucher_link"] = _voucher_link(request_id)
        return render_template(template, state="already", **page)
    if request.method == "GET":
        return render_template(template, state="pending", **page)

    direct = not r.get("ManagerCode")  # no reporting manager (HOD / senior staff): booking confirmed = approved
    rows, options, errors = parse_booking_options(request.form, allow_multiple=not direct)
    page.update(rows=rows, errors=errors)
    if errors:
        return render_template(template, state="pending", **page), 400

    to_status = STATUS_FINAL if direct else STATUS_MANAGER
    if not save_booking(request_id, options, to_status, comment=DIRECT_NOTE if direct else None):
        page["r"] = load_request(request_id)
        return render_template(template, state="already", **page)
    r = page["r"] = load_request(request_id)  # now includes the booking details

    if direct:
        page["voucher_link"] = _voucher_link(request_id)  # book now, or later from the email
        try:
            notify_approval(r)
        except Exception:  # noqa: BLE001
            log.exception("Final-booking email failed for %s", request_id)
        return render_template(template, state="done", **page)

    try:
        send_mail([r["ManagerEmail"]],
                  with_ref(r, f"Approval required: hotel booking for {r['EmpName']} ({r['EmpCode']})"),
                  manager_email_html(r))
    except Exception:  # noqa: BLE001
        log.exception("Failed to email the manager for %s", request_id)
        advance_status(request_id, STATUS_MANAGER, STATUS_BOOKING)  # details stay saved; the team can retry
        page["r"] = load_request(request_id)
        return render_template(template, state="pending", error=(
            "The booking was saved but could not be emailed to the manager. Please submit again."), **page), 502
    return render_template(template, state="done", **page)


@app.route("/review/<request_id>", methods=["GET", "POST"])
def review(request_id):
    """Step 2 - manager approves or rejects. GET only shows the page (safe for mail link scanners); POST decides.
    For a single choice, the email's Approve button opens a minimal confirm page (quick=1): the hotel and cost,
    the non-refundable tick-box and one Confirm Approval button."""
    token, err = _check_link(_review_signer, request_id, "review.html")
    if err:
        return err
    r = load_request(request_id)  # trust the DB record only
    if not r:
        abort(404)
    opts = booking_options(r)
    page = dict(r=r, token=token, fmt_date=fmt_date, booking_rows=booking_rows, logo_url=LOGO_URL)

    if request.method == "GET":
        if r["Status"] != STATUS_MANAGER:
            return render_template("review.html", state="already", **page)
        decision = request.args.get("decision", "")
        quick = decision == "approve" and request.args.get("quick") == "1" and len(opts) == 1
        return render_template("review.html", state="quick" if quick else "pending", decision=decision,
                               comment="", picked="", **page)

    decision = request.form.get("decision", "").lower()
    comment = request.form.get("comment", "").strip()[:300] if decision == "reject" else ""
    picked = request.form.get("option", "")
    acknowledged = request.form.get("ack") == "yes"
    page.update(picked=picked)
    if decision not in ("approve", "reject"):
        return render_template("review.html", state="pending", decision="", comment=comment,
                               error="Choose Approve or Reject.", **page), 400
    if decision == "reject" and not comment:
        return render_template("review.html", state="pending", decision="reject", comment="",
                               error="A reason is required to reject this request.", **page), 400

    option = option_no = None
    if decision == "approve":
        if len(opts) == 1 and not picked:
            picked = "1"  # a single choice is pre-selected
        if opts:
            if not (picked.isdigit() and 1 <= int(picked) <= len(opts)):
                return render_template("review.html", state="pending", decision="approve", comment="",
                                       error="Choose which hotel to approve.", **page), 400
            option_no = int(picked)
            option = opts[option_no - 1]
        if not acknowledged:
            quick = request.form.get("quick") == "1" and len(opts) == 1  # stay on the minimal confirm page
            return render_template("review.html", state="quick" if quick else "pending", decision="approve", comment="",
                                   error="Tick the box to acknowledge that the booking is non-refundable once "
                                         "approved.", **page), 400

    new_status = STATUS_FINAL if decision == "approve" else "Rejected"
    if not claim_decision(request_id, new_status, comment or None, option=option, option_no=option_no):
        page["r"] = load_request(request_id)
        return render_template("review.html", state="already", **page)

    r = page["r"] = load_request(request_id)  # includes the chosen choice for the emails
    try:
        if decision == "approve":
            notify_approval(r)
        else:
            notify_rejection(r, comment)
    except Exception:  # noqa: BLE001
        log.exception("Decision emails failed for %s", request_id)
    return render_template("review.html", state="done", **page)


def _safe_pdf_name(name, r):
    """A plain file name for the attachment, e.g. 'Voucher_Taj.pdf'; falls back to the booking ref."""
    base = re.sub(r"[^A-Za-z0-9 ._()-]+", "_", os.path.basename(name or "")).strip(" ._")[:120]
    if not base.lower().endswith(".pdf"):
        base = ""
    return base or f"Booking-voucher-{booking_ref(r).lstrip('#')}.pdf"


@app.route("/final-booking/<request_id>", methods=["GET", "POST"])
def final_booking(request_id):
    """Step 3 - after approval the Admin Team makes the final booking and uploads the voucher (PDF), which is
    emailed to the employee with the manager in CC. GET only shows the page; POST uploads and sends."""
    template = "final_booking.html"
    token, err = _check_link(_voucher_signer, request_id, template)
    if err:
        return err
    r = load_request(request_id)
    if not r:
        abort(404)
    page = dict(r=r, token=token, fmt_date=fmt_date, booking_rows=booking_rows, logo_url=LOGO_URL)

    if r["Status"] != STATUS_FINAL:
        return render_template(template, state="already", **page)
    if request.method == "GET":
        return render_template(template, state="pending", **page)

    upload = request.files.get("voucher")
    content = upload.read(MAX_VOUCHER_BYTES + 1) if upload else b""
    error = None
    if not upload or not upload.filename:
        error = "Choose the booking voucher (PDF) to upload."
    elif not content.startswith(b"%PDF-"):
        error = "The voucher must be a PDF file."
    elif len(content) > MAX_VOUCHER_BYTES:
        error = f"The voucher is too large. The maximum is {MAX_VOUCHER_BYTES // (1024 * 1024)} MB."
    if error:
        return render_template(template, state="pending", error=error, **page), 400

    file_name = _safe_pdf_name(upload.filename, r)
    if not save_voucher(request_id, file_name, content):
        page["r"] = load_request(request_id)
        return render_template(template, state="already", **page)
    r = page["r"] = load_request(request_id)

    cc = [r["ManagerEmail"]] if r.get("ManagerCode") and r.get("ManagerEmail") else None
    try:
        send_mail([r["EmpEmail"]],
                  with_ref(r, f"Your hotel booking in {r['PlaceTo']} is confirmed - voucher attached"),
                  voucher_email_html(r), cc=cc, attachments=[(file_name, "application/pdf", content)])
    except Exception:  # noqa: BLE001
        log.exception("Voucher email failed for %s", request_id)
        advance_status(request_id, STATUS_BOOKED, STATUS_FINAL)  # the team can upload again
        page["r"] = load_request(request_id)
        return render_template(template, state="pending", error=(
            "The voucher was saved but could not be emailed. Please upload it again."), **page), 502
    log.info("Voucher for %s emailed to the employee%s", booking_ref(r), " (manager in CC)" if cc else "")
    return render_template(template, state="done", file_name=file_name, **page)


@app.get("/final-booking/<request_id>/voucher")
def voucher_file(request_id):
    """The uploaded voucher, for the Admin Team (same signed link as the upload page)."""
    _, err = _check_link(_voucher_signer, request_id, "final_booking.html")
    if err:
        return err
    v = load_voucher(request_id)
    if not v:
        abort(404)
    name, content, _ = v
    return send_file(io.BytesIO(content), mimetype="application/pdf", download_name=name)


# --------------------------------------------------------------------------
# User Management - HR maintains the employee -> manager mapping that routes approval emails
# --------------------------------------------------------------------------
_CODE_RE = re.compile(r"^[A-Za-z0-9\-_/.]{1,50}$")
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_hr_failures = {}            # ip -> [timestamps of failed password attempts]
_hr_failures_lock = threading.Lock()
HR_MAX_FAILURES, HR_LOCKOUT_SECONDS = 5, 15 * 60


def _hr_networks():
    if HR_ALLOWED_NETWORKS.strip().lower() == "any":
        return None
    nets = []
    for part in HR_ALLOWED_NETWORKS.split(","):
        try:
            nets.append(ipaddress.ip_network(part.strip(), strict=False))
        except ValueError:
            log.warning("Ignoring invalid HR_ALLOWED_NETWORKS entry: %r", part)
    return nets


def hr_network_allowed(ip):
    if TRUSTED_PROXY and ip == TRUSTED_PROXY:
        # The proxy did not pass X-Forwarded-For, so the real visitor is unknown: refuse rather than trust it
        log.warning("User Management: no X-Forwarded-For from proxy %s; refusing", ip)
        return False
    nets = _hr_networks()
    if nets is None:
        return True
    try:
        addr = ipaddress.ip_address(ip)
    except ValueError:
        return False
    return any(addr in n for n in nets)


def _hr_locked_out(ip):
    now = time.time()
    with _hr_failures_lock:
        recent = [t for t in _hr_failures.get(ip, []) if now - t < HR_LOCKOUT_SECONDS]
        _hr_failures[ip] = recent
        return len(recent) >= HR_MAX_FAILURES


def _hr_record_failure(ip):
    with _hr_failures_lock:
        _hr_failures.setdefault(ip, []).append(time.time())


def hr_unlocked():
    """True while this browser's HR session is valid (sliding HR_SESSION_MINUTES idle timeout)."""
    last = session.get("hr_unlocked_at")
    if not last or time.time() - last > HR_SESSION_MINUTES * 60:
        session.pop("hr_unlocked_at", None)
        return False
    session["hr_unlocked_at"] = time.time()
    return True


def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_urlsafe(32)
    return session["csrf"]


def csrf_ok():
    sent = request.form.get("csrf", "")
    return bool(sent) and secrets.compare_digest(sent, session.get("csrf", ""))


app.jinja_env.globals["csrf_token"] = csrf_token


@app.before_request
def guard_user_management():
    if not request.path.startswith("/user-management"):
        return None
    if not hr_network_allowed(request.remote_addr or ""):
        log.warning("User Management blocked for %s (not in HR_ALLOWED_NETWORKS)", request.remote_addr)
        return render_template("hr_login.html", state="network", logo_url=LOGO_URL, show_hr_link=True), 403
    if not HR_ADMIN_PASSWORD:
        return render_template("hr_login.html", state="disabled", logo_url=LOGO_URL, show_hr_link=True), 503
    if request.endpoint in ("hr_login", "static") or hr_unlocked():
        return None
    if request.path.startswith("/user-management/api/"):
        return jsonify({"error": "locked"}), 401
    return redirect(url_for("hr_login"))


# ---- database
def list_mappings():
    """Every mapping row with names from the HR master (codes not in the master show no name)."""
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT m.id, m.EmpCode, e.EmployeeName, e.Status, {_q(EMP_EMAIL_COL)}, m.ManagerCode, mg.EmployeeName,
                   {_q(MGR_EMAIL_COL)}
            FROM {_q(MAPPING_TABLE)} m
            LEFT JOIN {_q(EMPLOYEE_TABLE)} e ON e.EmpCode = m.EmpCode
            LEFT JOIN {_q(EMPLOYEE_TABLE)} mg ON mg.EmpCode = m.ManagerCode
            ORDER BY m.EmpCode
            """
        )
        rows = cur.fetchall()
    return [{"id": int(r[0]), "emp_code": _clean(r[1]) or "", "emp_name": _clean(r[2]), "emp_status": _clean(r[3]),
             "emp_email": _clean(r[4]) or "", "mgr_code": _clean(r[5]) or "", "mgr_name": _clean(r[6]),
             "mgr_email": _clean(r[7]) or ""} for r in rows]


def get_mapping(mapping_id):
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT id, EmpCode, ManagerCode, {_q(EMP_EMAIL_COL)}, {_q(MGR_EMAIL_COL)} "
                    f"FROM {_q(MAPPING_TABLE)} WHERE id = ?", mapping_id)
        r = cur.fetchone()
    if not r:
        return None
    return {"id": int(r[0]), "emp_code": _clean(r[1]) or "", "mgr_code": _clean(r[2]) or "",
            "emp_email": _clean(r[3]) or "", "mgr_email": _clean(r[4]) or ""}


def mapping_id_for(emp_code):
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT TOP 1 id FROM {_q(MAPPING_TABLE)} WHERE EmpCode = ? ORDER BY id DESC", emp_code)
        r = cur.fetchone()
    return int(r[0]) if r else None


class SharedTableError(RuntimeError):
    """Raised instead of writing to a table that belongs to another application."""


def _writable_mapping_table():
    """The checklist apps own dbo.checklist_* tables; the portal must never write to them."""
    if "checklist_" in MAPPING_TABLE.lower():
        raise SharedTableError(f"{MAPPING_TABLE} is shared with the checklist apps; set MAPPING_TABLE to "
                               "dbo.TravelApproverMapping (migration 005)")
    return _q(MAPPING_TABLE)


def save_mapping(m, mapping_id=None):
    """Insert (mapping_id None) or update one mapping row in the portal's own table. Blank values become NULL."""
    table = _writable_mapping_table()
    vals = [m["emp_code"], m["mgr_code"] or None, m["emp_email"] or None, m["mgr_email"] or None]
    with get_conn() as conn:
        cur = conn.cursor()
        if mapping_id is None:
            cur.execute(f"INSERT INTO {table} (EmpCode, ManagerCode, {_q(EMP_EMAIL_COL)}, "
                        f"{_q(MGR_EMAIL_COL)}) VALUES (?, ?, ?, ?)", *vals)
            ok = True
        else:
            cur.execute(f"UPDATE {table} SET EmpCode = ?, ManagerCode = ?, {_q(EMP_EMAIL_COL)} = ?, "
                        f"{_q(MGR_EMAIL_COL)} = ?, UpdatedAt = SYSUTCDATETIME() WHERE id = ?", *vals, mapping_id)
            ok = cur.rowcount == 1
        conn.commit()
    return ok


def delete_mapping(mapping_id):
    table = _writable_mapping_table()
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"DELETE FROM {table} WHERE id = ?", mapping_id)
        ok = cur.rowcount == 1
        conn.commit()
    return ok


def master_lookup(code):
    """Name / email / status for a code from the HR master, or None."""
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT TOP 1 EmployeeName, EmployeeEmail, Status FROM {_q(EMPLOYEE_TABLE)} WHERE EmpCode = ?",
                    code)
        r = cur.fetchone()
    return {"name": _clean(r[0]), "email": _clean(r[1]), "status": _clean(r[2])} if r else None


def validate_mapping(form, mapping_id=None):
    """Returns (values, errors)."""
    v = {k: form.get(k, "").strip() for k in ("emp_code", "emp_email", "mgr_code", "mgr_email")}
    errors = {}
    if not _CODE_RE.match(v["emp_code"]):
        errors["emp_code"] = "Enter the employee code (letters, digits, - _ / . only)."
    if v["mgr_code"] and not _CODE_RE.match(v["mgr_code"]):
        errors["mgr_code"] = "Manager code can use letters, digits, - _ / . only."
    if v["mgr_code"] and v["mgr_code"].lower() == v["emp_code"].lower():
        errors["mgr_code"] = "An employee cannot be their own manager. Leave it blank for a HOD."
    for key in ("emp_email", "mgr_email"):
        if v[key] and (len(v[key]) > 256 or not _EMAIL_RE.match(v[key])):
            errors[key] = "Enter a valid email address, or leave it blank."
    if v["mgr_code"] and not v["mgr_email"]:
        errors["mgr_email"] = "The manager's email is needed to send approval requests."
    if "emp_code" not in errors:
        existing = mapping_id_for(v["emp_code"])
        if existing is not None and existing != mapping_id:
            errors["emp_code"] = f"{v['emp_code']} already has a mapping. Edit that row instead."
    return v, errors


def _audit(action, before=None, after=None):
    log.info("USER-MGMT %s from %s | before=%s | after=%s", action, request.remote_addr, before, after)


# ---- routes
@app.route("/user-management/login", methods=["GET", "POST"])
def hr_login():
    if hr_unlocked():
        return redirect(url_for("hr_dashboard"))
    ip = request.remote_addr or ""
    page = dict(logo_url=LOGO_URL, show_hr_link=True, hr_active=True)
    if _hr_locked_out(ip):
        return render_template("hr_login.html", state="locked", **page), 429
    if request.method == "GET":
        return render_template("hr_login.html", state="login", **page)
    if not csrf_ok():
        return render_template("hr_login.html", state="login", error="Your session expired. Please try again.",
                               **page), 400
    if secrets.compare_digest(request.form.get("password", "").encode(), HR_ADMIN_PASSWORD.encode()):
        session.clear()  # fresh session on unlock
        session["hr_unlocked_at"] = time.time()
        csrf_token()
        log.info("USER-MGMT unlocked from %s", ip)
        return redirect(url_for("hr_dashboard"))
    _hr_record_failure(ip)
    log.warning("USER-MGMT wrong password from %s", ip)
    if _hr_locked_out(ip):
        return render_template("hr_login.html", state="locked", **page), 429
    return render_template("hr_login.html", state="login", error="Incorrect password.", **page), 401


@app.post("/user-management/logout")
def hr_logout():
    if csrf_ok():
        session.pop("hr_unlocked_at", None)
    return redirect(url_for("hr_login"))


@app.get("/user-management")
def hr_dashboard(form=None, form_errors=None, status=200):
    try:
        rows = list_mappings()
    except Exception:  # noqa: BLE001
        log.exception("Could not load mappings")
        rows, form_errors = [], {"_": "Could not load the mappings from the database."}
    stats = {
        "total": len(rows),
        "no_manager": sum(1 for r in rows if not r["mgr_code"]),
        "missing_mgr_email": sum(1 for r in rows if r["mgr_code"] and not r["mgr_email"]),
        "missing_emp_email": sum(1 for r in rows if not r["emp_email"]),
    }
    return render_template("hr_users.html", rows=rows, stats=stats, form=form or {}, form_errors=form_errors or {},
                           logo_url=LOGO_URL, show_hr_link=True, hr_active=True,
                           session_minutes=HR_SESSION_MINUTES, this_month=date.today().strftime("%Y-%m"),
                           today=date.today().isoformat()), status


@app.post("/user-management/save")
def hr_save():
    if not csrf_ok():
        flash("Your session expired. Please try again.", "bad")
        return redirect(url_for("hr_dashboard"))
    raw_id = request.form.get("id", "").strip()
    mapping_id = int(raw_id) if raw_id.isdigit() else None
    before = get_mapping(mapping_id) if mapping_id else None
    if mapping_id and not before:
        flash("That mapping no longer exists.", "bad")
        return redirect(url_for("hr_dashboard"))
    values, errors = validate_mapping(request.form, mapping_id)
    if errors:
        return hr_dashboard(form={**values, "id": raw_id}, form_errors=errors, status=400)
    try:
        save_mapping(values, mapping_id)
    except SharedTableError:
        log.exception("Refused to write the mapping")
        flash("Saving is disabled: the portal is configured to use a table shared with other apps. "
              "Ask IT to set MAPPING_TABLE=dbo.TravelApproverMapping.", "bad")
        return redirect(url_for("hr_dashboard"))
    except pyodbc.IntegrityError:  # another HR user added the same Emp Code a moment ago
        return hr_dashboard(form={**values, "id": raw_id},
                            form_errors={"emp_code": f"{values['emp_code']} already has a mapping. Edit that row instead."},
                            status=400)
    _audit("UPDATE" if mapping_id else "ADD", before, values)
    who = values["emp_code"]
    flash(f"Mapping for {who} {'updated' if mapping_id else 'added'}. New hotel booking requests use it immediately.", "ok")
    return redirect(url_for("hr_dashboard", q=who))


@app.post("/user-management/delete")
def hr_delete():
    if not csrf_ok():
        flash("Your session expired. Please try again.", "bad")
        return redirect(url_for("hr_dashboard"))
    raw_id = request.form.get("id", "")
    before = get_mapping(int(raw_id)) if raw_id.isdigit() else None
    try:
        deleted = bool(before) and delete_mapping(before["id"])
    except SharedTableError:
        log.exception("Refused to delete from the mapping table")
        flash("Deleting is disabled: the portal is configured to use a table shared with other apps. "
              "Ask IT to set MAPPING_TABLE=dbo.TravelApproverMapping.", "bad")
        return redirect(url_for("hr_dashboard"))
    if not deleted:
        flash("That mapping no longer exists.", "bad")
        return redirect(url_for("hr_dashboard"))
    _audit("DELETE", before)
    flash(f"Mapping for {before['emp_code']} deleted.", "ok")
    return redirect(url_for("hr_dashboard"))


IST = timedelta(hours=5, minutes=30)  # the database stores times in UTC
EXPORT_COLUMNS = (  # (header, width)
    ("Booking Ref", 12), ("Status", 22), ("Requested on", 18), ("Emp Code", 11), ("Employee Name", 26),
    ("Designation", 22), ("Department", 20), ("Employee Email", 30), ("Manager", 24), ("Manager Email", 30),
    ("Travelling From", 18), ("Travelling To", 18), ("From (Date)", 13), ("To (Date)", 13), ("Nights", 8),
    ("Reason of Travel", 36), ("Category of Room", 18), ("Preferred Location", 24), ("Preferred Hotel", 24),
    ("Choices offered", 10), ("Approved choice", 10), ("Place/Area", 20), ("Hotel", 26), (COST_LABEL, 18),
    ("Type of occupancy", 18), ("Refund status", 18), ("Remarks", 30), ("Decided on", 18),
    ("Manager comment / Rejection reason", 36), ("Voucher emailed on", 18), ("Voucher file", 28),
)


def _ist(dt):
    return dt + IST if isinstance(dt, datetime) else None


def export_row(r):
    """One spreadsheet row per request, in EXPORT_COLUMNS order."""
    chosen = chosen_option(r) if is_approved(r) else {}
    return [
        r.get("BookingRefID"), r["Status"], _ist(r.get("CreatedAt")), r["EmpCode"], r["EmpName"],
        r.get("Designation"), r.get("Department"), r.get("EmpEmail"), r.get("ManagerName") or "No manager (HOD)",
        r.get("ManagerEmail") or None, r["PlaceFrom"], r["PlaceTo"], r["FromDate"], r["ToDate"],
        (r["ToDate"] - r["FromDate"]).days, r.get("ReasonOfTravel"), r.get("RoomCategory"),
        r.get("PreferredLocation"), r.get("PreferredHotel"), len(booking_options(r)) or None,
        selected_option_no(r) if is_approved(r) else None,
        r.get("BookingArea"), r.get("PropertyName"), r.get("Cost"),
        chosen.get("occupancy") or None, chosen.get("refund") or None, chosen.get("remarks") or None,
        _ist(r.get("ActionedAt")), r.get("ActionComment") or r.get("RejectionReason"),
        _ist(r.get("VoucherUploadedAt")), r.get("VoucherFileName"),
    ]


def build_export_xlsx(rows, start, end, basis):
    """The booking export as .xlsx bytes: a 'Bookings' sheet (filterable, frozen header) and an 'About' sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Bookings"
    ws.append([h for h, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="681313")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(export_row(r))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # text typed by users is never a formula
            elif isinstance(cell.value, datetime):
                cell.number_format = "dd-mmm-yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd-mmm-yyyy"
            elif isinstance(cell.value, Decimal):
                cell.number_format = "#,##0.00"
    for i, (_, width) in enumerate(EXPORT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(len(rows) + 1, 2)}"

    about = wb.create_sheet("About")
    for line in (
        ("Hotel booking export", None),
        ("Dates", f"{start:%d-%b-%Y} to {end:%d-%b-%Y}"),
        ("Filtered by", "Request date" if basis == "created" else "Check-in date (From)"),
        ("Requests", len(rows)),
        ("Generated", (datetime.utcnow() + IST).strftime("%d-%b-%Y %H:%M") + " (India time)"),
        ("Note", "Type of occupancy, Refund status and Remarks are Admin Team fields: do not share them with "
                 "employees or managers."),
    ):
        about.append(line)
    about["A1"].font = Font(bold=True, size=13)
    about.column_dimensions["A"].width = 16
    about.column_dimensions["B"].width = 70

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _export_range(args):
    """(start, end, label) from ?mode=month&month=YYYY-MM or ?mode=range&from=YYYY-MM-DD&to=YYYY-MM-DD."""
    if args.get("mode", "month") == "month":
        year, month = (int(x) for x in args.get("month", "").split("-"))
        start = date(year, month, 1)
        end = date(year, month, calendar.monthrange(year, month)[1])
        return start, end, f"{start:%b-%Y}"
    start, end = date.fromisoformat(args.get("from", "")), date.fromisoformat(args.get("to", ""))
    if end < start:
        raise ValueError("end before start")
    return start, end, f"{start:%d-%b-%Y}_to_{end:%d-%b-%Y}"


@app.get("/user-management/export")
def hr_export():
    """Download every hotel booking request in a month or date range as an Excel file."""
    basis = "checkin" if request.args.get("basis") == "checkin" else "created"
    try:
        start, end, label = _export_range(request.args)
    except ValueError:
        flash("Choose a month, or a From and To date (To on or after From), to export.", "bad")
        return redirect(url_for("hr_dashboard"))
    try:
        rows = export_requests(start, end, basis)
    except Exception:  # noqa: BLE001
        log.exception("Export query failed")
        flash("Could not read the bookings from the database. Please try again.", "bad")
        return redirect(url_for("hr_dashboard"))
    _audit("EXPORT", after=f"{start}..{end} by {basis}: {len(rows)} requests")
    return Response(build_export_xlsx(rows, start, end, basis),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Hotel-bookings_{label}.xlsx"'})


@app.get("/user-management/api/lookup/<code>")
def hr_lookup(code):
    """Name/email from the HR master while HR types a code in the form."""
    if not _CODE_RE.match(code):
        return jsonify({"found": False}), 400
    m = master_lookup(code)
    return jsonify({"found": bool(m), **(m or {})})


if __name__ == "__main__":
    from waitress import serve
    proxy = ({"trusted_proxy": TRUSTED_PROXY, "trusted_proxy_count": 1, "clear_untrusted_proxy_headers": True,
              "trusted_proxy_headers": {"x-forwarded-for", "x-forwarded-proto", "x-forwarded-host"}}
             if TRUSTED_PROXY else {})
    print(f"Serving Travel Accommodation Portal on http://{HOST}:{PORT}{APP_PREFIX}/ "
          f"(links use {APP_BASE_URL or 'APP_BASE_URL not set'}"
          f"{', behind proxy ' + TRUSTED_PROXY if TRUSTED_PROXY else ''})", flush=True)
    serve(app, host=HOST, port=PORT, threads=8, **proxy)
